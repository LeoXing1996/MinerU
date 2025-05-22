import numpy as np
import os
import copy
from datetime import datetime
from svgpathtools import svg2paths2
from scipy.optimize import minimize
import xml.etree.ElementTree as ET
from collections import Counter
from lxml import etree as lxml_etree
import shutil
import re
from pptx.dml.color import RGBColor
from pptx.util import Pt

from .shapes import *

from .shapes.util import (
    sample_svg_path,
    get_bounding_line,
    get_overall_bounding_rectangle,
    get_svg_paths,
    get_path_d,
    get_path_transform,
    get_path_attributes,
    parse_transform,
    apply_transform,
    judge_arrow_base,
    is_path_closed,
)
from .shapes.shape_pool import ShapePool
from .shapes.arrow import parse_arrows

# 全局形状计数器
shape_counter = Counter()


class ShapeOptimizer:
    def __init__(self):
        self.shape_pool = ShapePool()

    def error_function(self, params, target_points, shape_func):
        """
        计算模型生成的边界采样点与目标点集合之间的误差。
        对于模型中的每个点，取其到目标点集合中最近点的距离的平方，累计求和。
        """
        model_points = shape_func(params)
        error = 0
        for mp in model_points:
            # 计算当前模型点与所有目标点的欧氏距离
            d = np.linalg.norm(target_points - mp, axis=1)
            error += np.min(d) ** 2
        return error

    def optimize_shape(
        self,
        target_points,
        shape_name,
        shape_func,
        initial_guess,
        bounds=None,
        constraints=None,
        normalize=False,
    ):
        """
        针对给定候选形状，利用 Nelder-Mead 算法进行参数优化。
        输出形状名称、优化后参数和误差。
        """
        if len(initial_guess) > 0:
            options = dict(disp=False)  # print error value and iteration
            res = minimize(
                self.error_function,
                initial_guess,
                args=(target_points, shape_func),
                method='Nelder-Mead',
                tol=1e-6,  # 1e-5,
                bounds=bounds,
                constraints=constraints,
                options=options,
            )
            params = res.x
            err = res.fun
            # print(f"形状：{shape_name}\n优化后参数：{params}\n误差：{err}")
        else:
            params = initial_guess
            err = self.error_function(initial_guess, target_points, shape_func)
            # print(f"形状：{shape_name}\n优化后参数：{params}\n误差：{err}")
        return params, err

    def process_path(
        self, svg_path, num_points=200, normalize=False, error_threshold=50
    ):
        try:
            # get path and transform
            d = get_path_d(svg_path)

            if not d:  # 如果路径为空，直接返回None
                print('Empty path found, skipping...')
                return None, None

            # 如果路径不闭合，则不参与形状匹配
            if not (is_path_closed(d) or d[-1].lower() == 'z'):
                return None, -1

            # NOTE: 暂时不处理arrow path, 0426, 全部用svg path 来绘制
            # 判断一下是否是闭合的arrow path
            if judge_arrow_base(d):
                svg_path.attrib['d'] = judge_arrow_base(d)
                optimal_shape = self.process_arrow_base(svg_path)
                return optimal_shape, 0

            T_str = get_path_transform(svg_path)
            transform_matrix = parse_transform(T_str)

            # sample svg path
            target_points = sample_svg_path(
                d, num_points=num_points, normalize=normalize
            )
            target_points = apply_transform(target_points, transform_matrix)
            candidates = self.shape_pool.get_candidates(d)

            # TODO: lines, arrows, etc.
            if len(candidates) == 0:
                print('No candidates are found. Skip.')
                return None, None

            print('Candidates:', [cand.NAME for cand in candidates])

            min_error = 1e9
            optimal_shape = None

            for shape in candidates:
                for rotation, flip_h, flip_v in shape.TRANSFORMS:
                    s = shape()
                    s.set_definited_params(
                        target_points,
                        rotation=rotation,
                        flip_h=flip_h,
                        flip_v=flip_v,
                        normalize=normalize,
                    )
                    init_params = s.set_adjustable_params()

                    bounds = s.set_bounds()
                    constraints = s.set_constraints()

                    params, error = self.optimize_shape(
                        target_points,
                        s.NAME,
                        s.generate,
                        init_params,
                        bounds,
                        constraints,
                        normalize=normalize,
                    )

                    s.set_params(params)
                    s.set_transform(rotation, flip_h, flip_v)

                    if error < min_error:
                        min_error = min(error, min_error)
                        optimal_shape = s

            temp_svg_path = copy.deepcopy(svg_path)
            attributes = get_path_attributes(temp_svg_path)
            optimal_shape.set_attributes(attributes)
            points = optimal_shape.generate(optimal_shape.adjustments)

            if min_error >= error_threshold:
                return None, -2

            return optimal_shape, error if optimal_shape is not None else [None, None]
        except (ValueError, AttributeError) as e:
            print(f'Error processing path: {str(e)}')
            raise e
            return None, None

    def process_line(self, svg_path, arrow_end=False, num_points=200, normalize=False):
        # get path and transform
        d = get_path_d(svg_path)
        T_str = get_path_transform(svg_path)
        transform_matrix = parse_transform(T_str)

        # sample svg path
        target_points = sample_svg_path(d, num_points=num_points, normalize=normalize)
        target_points = apply_transform(target_points, transform_matrix)

        if arrow_end:
            arrow_bar = ArrowBar()
        else:
            arrow_bar = LineConnector()

        cx, cy, w, h = get_bounding_line(target_points)
        arrow_bar.w = w
        arrow_bar.h = h
        arrow_bar.cx = cx
        arrow_bar.cy = cy
        arrow_bar.theta = 0
        arrow_bar.rotation = 0
        arrow_bar.flip_h = False
        arrow_bar.flip_v = False

        arrow_bar.set_attributes(get_path_attributes(svg_path))
        _ = arrow_bar.generate(arrow_bar.adjustments)
        return arrow_bar

    def process_arrow_base(self, svg_path, num_points=200, normalize=False):
        # get path and transform
        d = get_path_d(svg_path)
        T_str = get_path_transform(svg_path)
        transform_matrix = parse_transform(T_str)

        # sample svg path
        target_points = sample_svg_path(d, num_points=num_points, normalize=normalize)
        target_points = apply_transform(target_points, transform_matrix)

        arrow_bar = ArrowBar()
        cx, cy, w, h = get_bounding_line(target_points)
        arrow_bar.w = w
        arrow_bar.h = h
        arrow_bar.cx = cx
        arrow_bar.cy = cy
        arrow_bar.theta = 0
        arrow_bar.flip_h = False
        arrow_bar.flip_v = False

        arrow_bar.set_attributes(get_path_attributes(svg_path))
        arrow_bar.set_transform(0, False, False)
        _ = arrow_bar.generate(arrow_bar.adjustments)

        arrow_bar.attributes['stroke'] = arrow_bar.attributes.get('fill', '#000000')
        arrow_bar.attributes['fill'] = 'None'
        arrow_bar.attributes['stroke-width'] = '1.0'

        return arrow_bar

    def convert_lxml_to_etree(self, lxml_element):
        """将lxml元素转换为xml.etree.ElementTree元素"""
        # 将lxml.etree._Attrib转换为dict
        attrib_dict = dict(lxml_element.attrib)
        etree_element = ET.Element(lxml_element.tag, attrib_dict)
        for child in lxml_element:
            etree_element.append(self.convert_lxml_to_etree(child))
        return etree_element

    def process_file(
        self,
        svg_file,
        num_points=200,
        normalize=False,
        prs=None,
        slide=None,
        records_dir=None,
        illustration_info=None,
    ):
        paths = get_svg_paths(svg_file)
        # print(paths)
        bbox = get_overall_bounding_rectangle(paths)

        # 用于收集未匹配成功的形状
        unmatched_paths = []

        arrow_paths = []
        for idx, path in enumerate(paths):
            # print(f"Processing path {idx} ...")
            # print("***" * 30)
            # print(path_to_svg(path))

            d = get_path_d(path)
            optimal_shape, error = self.process_path(
                path, num_points=num_points, normalize=normalize
            )

            ###log for debug #######################################
            # 如果找到了最优形状，记录到日志中
            if optimal_shape is not None:
                shape_counter[optimal_shape.NAME] += 1
                # ##############################################
                # # 匹配了形状，但是超过了threshold
                # if error==-2:
                #     unmatched_paths.append(path)
                # elif error!=-1:
                #     # 更新形状统计
                #     shape_counter[optimal_shape.NAME] += 1

                # ##############################################
                # log for debug
                # 确保日志目录存在
                log_dir = f'{records_dir}/log'
                if not os.path.exists(log_dir):
                    os.makedirs(log_dir)

                # 打开日志文件，追加模式写入
                base_name = os.path.splitext(os.path.basename(svg_file))[0]
                paper_name = svg_file.split('/')[1]
                log_file = os.path.join(log_dir, f'{paper_name}_{base_name}_log.txt')
                with open(log_file, 'a') as f:
                    timestamp = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                    f.write(f'--- Log Entry: {timestamp} ---\n')
                    f.write(f'Processing SVG Path: {d}\n')
                    f.write(f'Error: {error}\n')
                with open(log_file, 'a') as f:
                    f.write('Optimal Shape Found:\n')
                    f.write(f'  Shape Name: {optimal_shape.NAME}\n')
                    if optimal_shape.NAME != 'ARROW_BAR':
                        f.write(f'  Parameters: {optimal_shape.adjustments}\n')
                        f.write(f'  Drawing attributes: {optimal_shape.attributes}\n')
                        f.write(
                            f'  Transform: Rotation={optimal_shape.rotation}, Flip_H={optimal_shape.flip_h}, Flip_V={optimal_shape.flip_v}\n'
                        )
                        f.write('\n')
            #########################################################

            if optimal_shape is not None:
                prs, slide = optimal_shape.draw(prs, slide, save_name=None)
                # illustration_info['shape_code'].append(optimal_shape.to_dict())
            if optimal_shape is None and error == -1:  # path未闭合，error返回-1
                arrow_paths.append(path)
            if optimal_shape is None and (
                error == -2 or error is None
            ):  # 匹配了形状，但是超过了threshold，返回-2
                unmatched_paths.append(path)

        unmatched_paths.extend(arrow_paths)

        """
        # 暂时注释
        arrows, lines, others, pairs = parse_arrows(arrow_paths, bbox)
        unmatched_paths.extend(others)
        line_ids = {line_id for _, line_id in pairs}
        for idx, line in enumerate(lines):
            # print(f"Processing line {idx} ...")
            # print("***" * 30)
            # print(path_to_svg(line))
            if idx in line_ids:
                optimal_shape = self.process_line(line, arrow_end=True)
            else:
                optimal_shape = self.process_line(line)

            if optimal_shape is not None:
                prs, slide = optimal_shape.draw(prs, slide, save_name=None)
                # illustration_info['shape_code'].append(optimal_shape.to_dict())
                # 更新形状统计
                shape_counter[optimal_shape.NAME] += 1
            else:
                unmatched_paths.append(line)

            ###log for debug #######################################
            if optimal_shape is not None:
                log_dir = f"{records_dir}/log"
                if not os.path.exists(log_dir):
                    os.makedirs(log_dir)

                base_name = os.path.splitext(os.path.basename(svg_file))[0]
                paper_name = svg_file.split("/")[1]
                log_file = os.path.join(log_dir, f"{paper_name}_{base_name}_log.txt")
                with open(log_file, "a") as f:
                    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                    f.write(f"--- Log Entry: {timestamp} ---\n")
                    f.write(f"Processing SVG Path: {d}\n")
                with open(log_file, "a") as f:
                    f.write(f"Optimal Shape Found:\n")
                    f.write(f"  Shape Name: {optimal_shape.NAME}\n")
                    f.write(f"  Parameters: {optimal_shape.adjustments}\n")
                    f.write(f"  Drawing attributes: {optimal_shape.attributes}\n")
                    f.write("\n")
            #########################################################
        """

        # for idx, other in enumerate(others):

        #     print(f"Processing other {idx} ...")
        #     print("***" * 30)
        #     print(path_to_svg(other))
        #     other_d = get_path_d(other)

        #     flag = other.attrib.get("arrow_base", None)
        #     if flag:
        #         optimal_shape = self.process_arrow_base(other)
        #     else:
        #         optimal_shape = None

        #     if optimal_shape is not None:
        #         prs, slide = optimal_shape.draw(prs, slide, save_name=None)
        #         # 更新形状统计
        #         shape_counter[optimal_shape.NAME] += 1
        #     else:
        #         unmatched_paths.append(other)

        #     ###log for debug #######################################
        #     if optimal_shape is not None:
        #         log_dir = f"{records_dir}/log"
        #         if not os.path.exists(log_dir):
        #             os.makedirs(log_dir)

        #         base_name = os.path.splitext(os.path.basename(svg_file))[0]
        #         paper_name = svg_file.split("/")[1]
        #         log_file = os.path.join(log_dir, f"{paper_name}_{base_name}_log.txt")
        #         with open(log_file, "a") as f:
        #             timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        #             f.write(f"--- Log Entry: {timestamp} ---\n")
        #             f.write(f"Processing SVG Path: {d}\n")
        #         with open(log_file, "a") as f:
        #             f.write(f"Optimal Shape Found:\n")
        #             f.write(f"  Shape Name: {optimal_shape.NAME}\n")
        #             f.write(f"  Parameters: {optimal_shape.adjustments}\n")
        #             f.write(f"  Drawing attributes: {optimal_shape.attributes}\n")
        #             f.write("\n")
        #     #########################################################

        if unmatched_paths:
            # 创建保存未匹配形状的目录
            unmatched_dir = f'{records_dir}/unmatched_shapes'
            if not os.path.exists(unmatched_dir):
                os.makedirs(unmatched_dir)

            # 生成新的SVG文件名
            base_name = os.path.splitext(os.path.basename(svg_file))[0]
            paper_name = svg_file.split('/')[-4]
            os.makedirs(os.path.join(unmatched_dir, paper_name), exist_ok=True)
            unmatched_svg_path = os.path.join(
                unmatched_dir, paper_name, f'{base_name}_unmatched.svg'
            )
            original_svg_path = os.path.join(
                unmatched_dir, paper_name, f'{base_name}_original.svg'
            )

            # 复制原始SVG文件
            shutil.copy2(svg_file, original_svg_path)
            print(f'保存原始SVG文件到: {original_svg_path}')

            # 读取SVG文件并创建树
            with open(svg_file, 'r') as file:
                svg_content = file.read()
            svg_tree = lxml_etree.fromstring(svg_content)
            namespaces = {'svg': 'http://www.w3.org/2000/svg'}

            # 获取所有路径元素
            all_paths = svg_tree.xpath('//svg:path', namespaces=namespaces)
            all_paths = [
                path
                for path in all_paths
                if path.getparent().tag != f"{{{namespaces['svg']}}}clipPath"
            ]

            # 删除所有已匹配的路径
            for path in all_paths:
                # 检查当前路径是否在未匹配路径列表中
                is_matched = False
                for unmatched_path in unmatched_paths:
                    if path.get('d') == unmatched_path.get('d'):
                        is_matched = True
                        break

                if not is_matched:
                    # 找到路径的父元素并删除路径
                    parent = path.getparent()
                    if parent is not None:
                        parent.remove(path)

            # 保存修改后的SVG
            lxml_etree.ElementTree(svg_tree).write(
                unmatched_svg_path,
                encoding='utf-8',
                xml_declaration=False,
                pretty_print=True,
            )
            print(f'保存未匹配的形状到: {unmatched_svg_path}')

            #########################################################
            # 利用picosvg简化paths
            from picosvg.svg import SVG
            from typing import Union, Optional

            def simplify_svg(
                input_svg: Union[str, bytes],
                output_path: Optional[str] = None,
            ) -> str:
                """
                使用 picosvg 简化 SVG 路径

                Args:
                    input_svg: SVG 文件路径或 SVG 内容字符串
                    output_path: 输出 SVG 文件路径，如不提供则只返回结果不保存

                Returns:
                    str: 简化后的 SVG 内容

                Example:
                    >>> # 简化 SVG 文件并直接返回内容
                    >>> simplified_svg = simplify_svg("complex.svg")
                    >>> # 简化 SVG 文件并保存
                    >>> simplify_svg("complex.svg", "simplified.svg")
                    >>> # 从字符串简化
                    >>> svg_content = '<svg>...</svg>'
                    >>> simplified_svg = simplify_svg(svg_content)
                """
                try:
                    # 判断输入是文件路径还是 SVG 内容
                    if isinstance(input_svg, str) and os.path.isfile(input_svg):
                        svg = SVG.parse(input_svg)
                    else:
                        # 假设输入是 SVG 内容字符串
                        svg = SVG.fromstring(input_svg)

                    # 简化 SVG
                    svg = svg.topicosvg()

                    # 生成结果
                    simplified_svg = svg.tostring()

                    # 如果提供了输出路径，则保存文件
                    if output_path:
                        with open(output_path, 'w', encoding='utf-8') as f:
                            f.write(simplified_svg)
                        print(f'已保存简化的 SVG 到 {output_path}')

                    return simplified_svg

                except Exception as e:
                    print(f'简化 SVG 时发生错误: {str(e)}')
                    raise e

            simplified_unmatched_svg_path = os.path.join(
                unmatched_dir, paper_name, f'{base_name}_simplified_unmatched.svg'
            )
            # simplified_svg = simplify_svg(
            #     unmatched_svg_path, simplified_unmatched_svg_path
            # )
            # slide = add_svg_to_slide(slide, simplified_unmatched_svg_path)
            #########################################################

        """
        # 生成形状统计报告
        stats_dir = "shape_statistics"
        os.makedirs(stats_dir, exist_ok=True)

        # 写入形状统计
        with open(os.path.join(stats_dir, "matched_shapes.csv"), 'a', newline='') as f:
            writer = csv.writer(f)
            writer.writerow(["Shape Type", "Count"])
            for shape, count in shape_counter.most_common():
                writer.writerow([shape, count])

        # 生成总统计报告
        with open(os.path.join(stats_dir, "shape_statistics_report.txt"), 'a') as f:
            f.write("形状类型统计报告\n")
            f.write("=" * 50 + "\n\n")

            f.write("已匹配形状统计:\n")
            f.write("-" * 30 + "\n")
            total = sum(shape_counter.values())
            for shape, count in shape_counter.most_common():
                percentage = count / total * 100 if total > 0 else 0
                f.write(f"{shape}: {count} ({percentage:.2f}%)\n")

            f.write("\n总计:\n")
            f.write("-" * 30 + "\n")
            f.write(f"形状总数: {total}\n")
        #########################################################
        """
        return prs, slide

    def convert_svg_to_illustration(
        self, svg_file, num_points=200, normalize=False, records_dir=None
    ):
        paths = get_svg_paths(svg_file)
        bbox = get_overall_bounding_rectangle(paths)

        # 用于收集未匹配成功的形状
        unmatched_paths = []
        shape_illustration = []
        path_illustration = []
        arrow_paths = []
        for idx, path in enumerate(paths):
            # print(f"Processing path {idx} ...")
            # print("***" * 30)
            # print(path_to_svg(path))

            d = get_path_d(path)
            optimal_shape, error = self.process_path(
                path, num_points=num_points, normalize=normalize
            )

            ###log for debug #######################################
            # 如果找到了最优形状，记录到日志中
            if optimal_shape is not None:
                shape_counter[optimal_shape.NAME] += 1
                # ##############################################
                # # 匹配了形状，但是超过了threshold
                # if error==-2:
                #     unmatched_paths.append(path)
                # elif error!=-1:
                #     # 更新形状统计
                #     shape_counter[optimal_shape.NAME] += 1

            if optimal_shape is not None:
                shape_illustration.append(optimal_shape.to_dict())
            if optimal_shape is None and error == -1:  # path未闭合，error返回-1
                arrow_paths.append(path)
            if optimal_shape is None and (
                error == -2 or error is None
            ):  # 匹配了形状，但是超过了threshold，返回-2
                unmatched_paths.append(path)

        # unmatched_paths.extend(arrow_paths)

        arrows, lines, others, pairs = parse_arrows(arrow_paths, bbox)
        unmatched_paths.extend(others)
        line_ids = {line_id for _, line_id in pairs}
        for idx, line in enumerate(lines):
            # print(f"Processing line {idx} ...")
            # print("***" * 30)
            # print(path_to_svg(line))
            if idx in line_ids:
                optimal_shape = self.process_line(line, arrow_end=True)
            else:
                optimal_shape = self.process_line(line)

        """
        # 暂时注释
        arrows, lines, others, pairs = parse_arrows(arrow_paths, bbox)
        unmatched_paths.extend(others)
        line_ids = {line_id for _, line_id in pairs}
        for idx, line in enumerate(lines):
            # print(f"Processing line {idx} ...")
            # print("***" * 30)
            # print(path_to_svg(line))
            if idx in line_ids:
                optimal_shape = self.process_line(line, arrow_end=True)
            else:
                optimal_shape = self.process_line(line)

            if optimal_shape is not None:
                prs, slide = optimal_shape.draw(prs, slide, save_name=None)
                # illustration_info['shape_code'].append(optimal_shape.to_dict())
                # 更新形状统计
                shape_counter[optimal_shape.NAME] += 1
            else:
                unmatched_paths.append(line)

            ###log for debug #######################################
            if optimal_shape is not None:
                log_dir = f"{records_dir}/log"
                if not os.path.exists(log_dir):
                    os.makedirs(log_dir)

                base_name = os.path.splitext(os.path.basename(svg_file))[0]
                paper_name = svg_file.split("/")[1]
                log_file = os.path.join(log_dir, f"{paper_name}_{base_name}_log.txt")
                with open(log_file, "a") as f:
                    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                    f.write(f"--- Log Entry: {timestamp} ---\n")
                    f.write(f"Processing SVG Path: {d}\n")
                with open(log_file, "a") as f:
                    f.write(f"Optimal Shape Found:\n")
                    f.write(f"  Shape Name: {optimal_shape.NAME}\n")
                    f.write(f"  Parameters: {optimal_shape.adjustments}\n")
                    f.write(f"  Drawing attributes: {optimal_shape.attributes}\n")
                    f.write("\n")
            #########################################################
        """

        if unmatched_paths:
            # 创建保存未匹配形状的目录
            unmatched_dir = f'{records_dir}/unmatched_shapes'
            if not os.path.exists(unmatched_dir):
                os.makedirs(unmatched_dir)

            # 生成新的SVG文件名
            base_name = os.path.splitext(os.path.basename(svg_file))[0]
            paper_name = svg_file.split('/')[-4]
            os.makedirs(os.path.join(unmatched_dir, paper_name), exist_ok=True)
            unmatched_svg_path = os.path.join(
                unmatched_dir, paper_name, f'{base_name}_unmatched.svg'
            )
            original_svg_path = os.path.join(
                unmatched_dir, paper_name, f'{base_name}_original.svg'
            )

            # 复制原始SVG文件
            shutil.copy2(svg_file, original_svg_path)
            print(f'保存原始SVG文件到: {original_svg_path}')

            # 读取SVG文件并创建树
            with open(svg_file, 'r') as file:
                svg_content = file.read()
            svg_tree = lxml_etree.fromstring(svg_content)
            namespaces = {'svg': 'http://www.w3.org/2000/svg'}

            # 获取所有路径元素
            all_paths = svg_tree.xpath('//svg:path', namespaces=namespaces)
            all_paths = [
                path
                for path in all_paths
                if path.getparent().tag != f"{{{namespaces['svg']}}}clipPath"
            ]

            # 删除所有已匹配的路径
            for path in all_paths:
                # 检查当前路径是否在未匹配路径列表中
                is_matched = False
                for unmatched_path in unmatched_paths:
                    if path.get('d') == unmatched_path.get('d'):
                        is_matched = True
                        break

                if not is_matched:
                    # 找到路径的父元素并删除路径
                    parent = path.getparent()
                    if parent is not None:
                        parent.remove(path)

            # 保存修改后的SVG
            lxml_etree.ElementTree(svg_tree).write(
                unmatched_svg_path,
                encoding='utf-8',
                xml_declaration=False,
                pretty_print=True,
            )
            print(f'保存未匹配的形状到: {unmatched_svg_path}')

            #########################################################
            # 利用picosvg简化paths
            from picosvg.svg import SVG
            from typing import Union, Optional

            def simplify_svg(
                input_svg: Union[str, bytes],
                output_path: Optional[str] = None,
            ) -> str:
                """
                使用 picosvg 简化 SVG 路径

                Args:
                    input_svg: SVG 文件路径或 SVG 内容字符串
                    output_path: 输出 SVG 文件路径，如不提供则只返回结果不保存

                Returns:
                    str: 简化后的 SVG 内容

                Example:
                    >>> # 简化 SVG 文件并直接返回内容
                    >>> simplified_svg = simplify_svg("complex.svg")
                    >>> # 简化 SVG 文件并保存
                    >>> simplify_svg("complex.svg", "simplified.svg")
                    >>> # 从字符串简化
                    >>> svg_content = '<svg>...</svg>'
                    >>> simplified_svg = simplify_svg(svg_content)
                """
                try:
                    # 判断输入是文件路径还是 SVG 内容
                    if isinstance(input_svg, str) and os.path.isfile(input_svg):
                        svg = SVG.parse(input_svg)
                    else:
                        # 假设输入是 SVG 内容字符串
                        svg = SVG.fromstring(input_svg)

                    # 简化 SVG
                    svg = svg.topicosvg()

                    # 生成结果
                    simplified_svg = svg.tostring()

                    # 如果提供了输出路径，则保存文件
                    if output_path:
                        with open(output_path, 'w', encoding='utf-8') as f:
                            f.write(simplified_svg)
                        print(f'已保存简化的 SVG 到 {output_path}')

                    return simplified_svg

                except Exception as e:
                    print(f'简化 SVG 时发生错误: {str(e)}')
                    raise e

            simplified_unmatched_svg_path = os.path.join(
                unmatched_dir, paper_name, f'{base_name}_simplified_unmatched.svg'
            )
            simplified_svg = simplify_svg(
                unmatched_svg_path, simplified_unmatched_svg_path
            )
            # slide = add_svg_to_slide(slide, simplified_unmatched_svg_path)
            # TODO: 将简化后的svg添加到path_illustration中
            # path_illustration.append(simplified_unmatched_svg_path)

        return shape_illustration, simplified_unmatched_svg_path


def add_svg_to_slide(slide, svg_path, sample_points=20):
    """
    将未匹配的 SVG 路径直接添加到幻灯片中

    Args:
        slide: PowerPoint 幻灯片对象
        svg_path: SVG 路径
        sample_points: 路径采样点数量
    """

    def hex_to_rgb(hex_color):
        hex_color = hex_color.strip()
        if hex_color.startswith('#'):
            hex_color = hex_color.lstrip('#')
            return tuple(int(hex_color[i : i + 2], 16) for i in (0, 2, 4))
        elif hex_color.startswith('rgb'):
            nums = map(int, re.findall(r'\d+', hex_color))
            return tuple(nums)
        else:
            return (0, 0, 0)  # 默认黑色

    paths, attributes, svg_attributes = svg2paths2(svg_path)
    for path, attr in zip(paths, attributes):
        builder = slide.shapes.build_freeform()

        # 设置起始点
        start_point = path[0].start
        start_x, start_y = start_point.real, start_point.imag
        start_x, start_y = Pt(start_x), Pt(start_y)
        builder.move_to(start_x, start_y)

        # 采样路径点
        following_points = []
        for segment in path:
            points = [segment.point(t) for t in np.linspace(0, 1, sample_points)]
            for p in points:
                x, y = p.real, p.imag
                x, y = Pt(x), Pt(y)
                following_points.append((x, y))

        # 添加路径点并转换为形状
        builder.add_line_segments(following_points, close=True)
        shape = builder.convert_to_shape()

        # 设置填充
        shape.fill.solid()
        fill_color = attr.get('fill', '#000000')
        if fill_color != 'none':
            fill_color = hex_to_rgb(fill_color)
            shape.fill.fore_color.rgb = RGBColor(*fill_color)

        # 禁用阴影
        shape.shadow.inherit = False

        # 设置边框（无边框）
        shape.line.color.rgb = RGBColor(*fill_color)
        shape.line.width = Pt(0)
        shape.line.fill.background()

    return slide


if __name__ == '__main__':
    file_path = '/cpfs01/user/maqianli/LLM4Pipeline/MinerU/test/2412.15205v1/auto/images/page1_Fig2_0.svg'
    optimizer = ShapeOptimizer()
    optimizer.process_file(file_path)
